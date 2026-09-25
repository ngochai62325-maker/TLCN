"""World Bank Commodity Markets Pink Sheet Adapter.

Handles ingestion of World Bank historical commodity market price datasets
(CMO-Historical-Data-Monthly.xlsx).

Features:
- Pure LOCAL_FILE runtime operation (zero HTTP/network requests).
- Specialized multi-row header parsing for sheet 'Monthly Prices'
  (extracting commodity names from row 4 and units from row 5,
  preserving 'Period' date column).
- Pure raw-data semantics preservation without artificial melting/joins.
- Checksum SHA-256 computation and source file traceability (_source_file).
- Strict Ingestion Exception Hierarchy classification.
"""

from __future__ import annotations

import os
from typing import Any, Generator, List, Optional

import openpyxl
import pandas as pd

from ingestion.adapters.base_adapter import BaseSourceAdapter
from ingestion.core.config import SourceConfig
from ingestion.core.enums import SourceType
from ingestion.core.result import DataChunk, ReadinessResult
from ingestion.utils.error_classifier import (
    DataQualityError,
    PermanentError,
    SchemaError,
)
from ingestion.utils.hashing import compute_file_checksum
from ingestion.utils.logging_config import create_ingestion_logger


class WorldBankAdapter(BaseSourceAdapter):
    """Adapter for World Bank Pink Sheet commodity price data."""

    def check_readiness(self, config: SourceConfig) -> ReadinessResult:
        """Check availability and readiness of the local World Bank Excel file."""
        logger = create_ingestion_logger(config.source_id)

        target_file = config.local_path or config.local_fallback
        if not target_file:
            return ReadinessResult(ready=False, reason="Missing local_path for World Bank adapter")

        if not os.path.exists(target_file):
            return ReadinessResult(ready=False, reason=f"Local file not found: {target_file}")

        if not os.path.isfile(target_file):
            return ReadinessResult(ready=False, reason=f"Local path is not a file: {target_file}")

        if not os.access(target_file, os.R_OK):
            return ReadinessResult(ready=False, reason=f"Local file not readable: {target_file}")

        file_size = os.path.getsize(target_file)
        if file_size == 0:
            return ReadinessResult(ready=False, reason=f"Local file is empty (0 bytes): {target_file}")

        min_size = getattr(config.readiness, "min_file_size_bytes", 0) if config.readiness else 0
        if min_size and file_size < min_size:
            return ReadinessResult(
                ready=False,
                reason=f"File size {file_size} bytes is less than min_file_size_bytes {min_size}",
            )

        if not (target_file.lower().endswith(".xlsx") or target_file.lower().endswith(".xls")):
            return ReadinessResult(
                ready=False,
                reason=f"Unsupported file format: '{target_file}'. Expected .xlsx or .xls",
            )

        sheet_name = (config.extra.get("sheet_name") if config.extra else None) or "Monthly Prices"
        try:
            wb = openpyxl.load_workbook(target_file, read_only=True, data_only=True)
            available_sheets = wb.sheetnames
            wb.close()
            if sheet_name not in available_sheets:
                return ReadinessResult(
                    ready=False,
                    reason=f"Expected sheet '{sheet_name}' not found. Available sheets: {available_sheets}",
                )
        except Exception as e:
            return ReadinessResult(
                ready=False,
                reason=f"Failed to inspect Excel workbook '{target_file}': {e}",
            )

        return ReadinessResult(
            ready=True,
            reason=f"Local Excel file is ready with sheet '{sheet_name}' ({file_size:,} bytes)",
            source_metadata={
                "path": target_file,
                "size_bytes": file_size,
                "sheet_name": sheet_name,
                "available_sheets": available_sheets,
            },
        )

    def extract_full(
        self,
        config: SourceConfig,
        download_dir: str,
        resume_chunk_id: Optional[int] = None,
        resume_row_start: Optional[int] = None,
    ) -> Generator[DataChunk, None, None]:
        """Extract dataset from World Bank Pink Sheet Excel file."""
        logger = create_ingestion_logger(config.source_id)

        target_file = config.local_path or config.local_fallback
        if not target_file or not os.path.exists(target_file):
            raise FileNotFoundError(f"Local World Bank file not found: '{target_file}'")

        if not (target_file.lower().endswith(".xlsx") or target_file.lower().endswith(".xls")):
            raise PermanentError(f"Unsupported file format for World Bank adapter: '{target_file}'")

        file_size = os.path.getsize(target_file)
        if file_size == 0:
            logger.warning(f"World Bank Excel file '{target_file}' is empty (0 bytes).")
            return

        sheet_name = (config.extra.get("sheet_name") if config.extra else None) or "Monthly Prices"

        try:
            wb = openpyxl.load_workbook(target_file, read_only=True, data_only=True)
        except Exception as e:
            raise DataQualityError(f"Corrupted or invalid Excel workbook '{target_file}': {e}")

        if sheet_name not in wb.sheetnames:
            wb.close()
            raise DataQualityError(
                f"Expected sheet '{sheet_name}' not found in workbook '{target_file}'. Available: {wb.sheetnames}"
            )

        try:
            ws = wb[sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))
            wb.close()
        except Exception as e:
            wb.close()
            raise DataQualityError(f"Failed to read sheet '{sheet_name}' from '{target_file}': {e}")

        # World Bank Pink Sheet Layout:
        # Rows 0-3: Titles and preamble metadata
        # Row 4: Commodity names
        # Row 5: Units
        # Rows 6+: Monthly data observations (Col 0 = Period 'YYYYMmm')
        if len(all_rows) < 7:
            raise DataQualityError(
                f"Sheet '{sheet_name}' has insufficient rows ({len(all_rows)}). Expected at least 7 rows."
            )

        comm_row = all_rows[4]
        unit_row = all_rows[5]
        data_rows = all_rows[6:]

        col_names: List[str] = ["Period"]
        for idx in range(1, len(comm_row)):
            comm = str(comm_row[idx]).strip() if comm_row[idx] is not None else f"Commodity_{idx}"
            unit = str(unit_row[idx]).strip() if unit_row[idx] is not None else ""
            col_name = f"{comm} {unit}".strip() if unit else comm
            col_names.append(col_name)

        # Build DataFrame
        df = pd.DataFrame(data_rows, columns=col_names)

        # Drop any entirely empty trailing rows if present
        df.dropna(how="all", inplace=True)
        record_count = len(df)
        if record_count == 0:
            logger.warning(f"Sheet '{sheet_name}' in '{target_file}' contains no data records.")
            return

        # Technical Schema Contract verification
        if config.expected_columns:
            missing_cols = [c for c in config.expected_columns if c not in df.columns]
            if missing_cols:
                raise SchemaError(
                    f"Schema mismatch in '{target_file}': missing required columns: {missing_cols}"
                )

        # Attach traceability metadata
        df["_source_file"] = os.path.basename(target_file)

        file_checksum = compute_file_checksum(target_file)

        data_chunk = DataChunk(
            chunk_id=0,
            data=df,
            row_start=0,
            row_end=record_count - 1,
            record_count=record_count,
            checksum=file_checksum,
        )

        logger.info(
            f"Yielded DataChunk from '{target_file}' sheet '{sheet_name}': {record_count} records, {len(df.columns)} columns"
        )
        yield data_chunk

    def extract_incremental(
        self,
        config: SourceConfig,
        download_dir: str,
        watermark: Any,
    ) -> Generator[DataChunk, None, None]:
        """Incremental extraction not implemented for snapshot World Bank dataset."""
        raise NotImplementedError("Incremental extract not implemented for WorldBankAdapter")

    def get_artifact_checksum(self, file_path: str) -> str:
        """Compute SHA-256 checksum of file."""
        return compute_file_checksum(file_path)
