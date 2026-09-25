import os
import sys
import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding='utf-8')

def inspect_excel(path):
    print(f"\n=======================================================")
    print(f"FILE: {path} ({os.path.getsize(path):,} bytes)")
    print(f"=======================================================")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    print(f"Sheet names ({len(sheet_names)}): {sheet_names}")
    
    for sname in sheet_names:
        ws = wb[sname]
        print(f"\n--- SHEET: '{sname}' ---")
        # Read first 15 rows with openpyxl
        rows = list(ws.iter_rows(values_only=True, max_row=15))
        print(f"Total preview rows read: {len(rows)}")
        for idx, r in enumerate(rows[:10]):
            non_empty = [c for c in r if c is not None]
            print(f"  Row {idx}: non-empty={len(non_empty)}, sample={r[:6]}")
            
    wb.close()

inspect_excel("data/raw/world_bank/CMO-Historical-Data-Monthly.xlsx")
inspect_excel("data/raw/thitruongnongsan/price_luagao.xlsx")
